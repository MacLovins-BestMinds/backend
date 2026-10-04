import io

import pytest

from app.ai.slides import DEMO_TEXT, PDF, PPTX, FitSlide, FitSlidesDraft, SlidesError, assemble, file_kind, pptx_slides


def _deck() -> bytes:
    from pptx import Presentation

    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[1])
    first.shapes.title.text = "Smart pill box"
    first.placeholders[1].text = "Older people miss their medicine"
    demo = deck.slides.add_slide(deck.slide_layouts[5])
    demo.shapes.title.text = "Live demo"
    demo.notes_slide.notes_text_frame.text = "open the app"
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def test_pptx_slides_keep_titles_text_and_notes() -> None:
    data = _deck()
    assert file_kind("deck.pptx", data) == PPTX
    first, demo = pptx_slides(data)
    assert "Title: Smart pill box" in first and "Older people miss their medicine" in first
    assert "Title: Live demo" in demo and "Speaker notes: open the app" in demo


def test_file_kind_goes_by_content() -> None:
    assert file_kind("anything.bin", b"%PDF-1.7 ...") == PDF
    with pytest.raises(SlidesError):
        file_kind("notes.txt", b"just text")


def test_demo_slide_always_says_demo_time() -> None:
    draft = FitSlidesDraft(
        slides=[
            FitSlide(n=1, title="Problem", kind="talk", text=" People forget pills. "),
            FitSlide(n=7, title="Demo", kind="demo", text="Now I will show you how the app works in detail."),
        ]
    )
    result = assemble(draft)
    assert [(s.n, s.text) for s in result.slides] == [(1, "People forget pills."), (2, DEMO_TEXT)]
    assert result.text == "[Slide 1 — Problem]\nPeople forget pills.\n\n[Slide 2 — Demo]\nDemo time."
